
from pathlib import Path
from datetime import datetime

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.chart import BarChart, Reference
from openpyxl.utils import get_column_letter


# Project paths
ROOT = Path(__file__).resolve().parent.parent
DATA_DIR = ROOT / "Dataset" / "cleaned"
OUTPUT_DIR = ROOT / "Excel"
OUTPUT_FILE = OUTPUT_DIR / "Stock_Market_Dashboard.xlsx"

COMPANY_FILE = DATA_DIR / "company_summary.csv"
SECTOR_FILE = DATA_DIR / "sector_summary.csv"


def main():
    # Check that the summary files exist
    if not COMPANY_FILE.exists() or not SECTOR_FILE.exists():
        raise FileNotFoundError(
            "Summary CSV files not found. Run eda.py first."
        )

    companies = pd.read_csv(COMPANY_FILE)
    sectors = pd.read_csv(SECTOR_FILE)

    # Basic validation
    if companies.empty or sectors.empty:
        raise ValueError("A summary CSV file is empty.")

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    dashboard = wb.active
    dashboard.title = "Dashboard"

    company_ws = wb.create_sheet("Company Analysis")
    sector_ws = wb.create_sheet("Sector Analysis")

    # Theme
    navy = "17365D"
    blue = "D9EAF7"
    white = "FFFFFF"
    green = "E2F0D9"

    # Dashboard heading
    dashboard.merge_cells("A1:H1")
    dashboard["A1"] = "NSE STOCK MARKET ANALYSIS DASHBOARD"
    dashboard["A1"].font = Font(
        bold=True, size=16, color=white
    )
    dashboard["A1"].fill = PatternFill(
        "solid", fgColor=navy
    )
    dashboard["A1"].alignment = Alignment(
        horizontal="center"
    )
    dashboard.row_dimensions[1].height = 32

    dashboard["A3"] = "Last generated"
    dashboard["B3"] = datetime.now().strftime(
        "%Y-%m-%d %H:%M"
    )

    # KPI cards
    kpis = [
        ("Companies analyzed", len(companies)),
        ("Sectors analyzed", len(sectors)),
        ("Best company return (%)",
         round(companies["Return"].max(), 2)),
        ("Lowest company return (%)",
         round(companies["Return"].min(), 2)),
    ]

    for col, (label, value) in enumerate(kpis, start=1):
        title_cell = dashboard.cell(row=5, column=col, value=label)
        value_cell = dashboard.cell(row=6, column=col, value=value)

        title_cell.font = Font(bold=True, color=white)
        title_cell.fill = PatternFill("solid", fgColor=navy)
        title_cell.alignment = Alignment(
            horizontal="center", wrap_text=True
        )

        value_cell.font = Font(bold=True, size=14)
        value_cell.fill = PatternFill("solid", fgColor=green)
        value_cell.alignment = Alignment(horizontal="center")

    # Write company analysis table
    for row in [list(companies.columns)] + companies.values.tolist():
        company_ws.append(row)

    # Write sector analysis table
    for row in [list(sectors.columns)] + sectors.values.tolist():
        sector_ws.append(row)

    # Format table headers and column widths
    for ws in [company_ws, sector_ws]:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

        for cell in ws[1]:
            cell.font = Font(bold=True, color=white)
            cell.fill = PatternFill("solid", fgColor=navy)
            cell.alignment = Alignment(
                horizontal="center", wrap_text=True
            )

        for column in ws.columns:
            letter = get_column_letter(column[0].column)
            max_length = max(
                len(str(cell.value or "")) for cell in column
            )
            ws.column_dimensions[letter].width = min(
                max_length + 3, 32
            )

    # Company returns chart
    chart = BarChart()
    chart.type = "bar"
    chart.style = 10
    chart.title = "Company Period Returns"
    chart.x_axis.title = "Return (%)"
    chart.y_axis.title = "Company"
    chart.height = 9
    chart.width = 17

    return_col = list(companies.columns).index("Return") + 1
    data = Reference(
        company_ws,
        min_col=return_col,
        min_row=1,
        max_row=len(companies) + 1,
    )
    categories = Reference(
        company_ws,
        min_col=1,
        min_row=2,
        max_row=len(companies) + 1,
    )
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(categories)
    dashboard.add_chart(chart, "A9")

    # Sector returns chart
    sector_chart = BarChart()
    sector_chart.type = "bar"
    sector_chart.style = 11
    sector_chart.title = "Average Return by Sector"
    sector_chart.x_axis.title = "Return (%)"
    sector_chart.y_axis.title = "Sector"
    sector_chart.height = 9
    sector_chart.width = 17

    sector_return_col = (
        list(sectors.columns).index("Average Return") + 1
    )
    sector_data = Reference(
        sector_ws,
        min_col=sector_return_col,
        min_row=1,
        max_row=len(sectors) + 1,
    )
    sector_categories = Reference(
        sector_ws,
        min_col=1,
        min_row=2,
        max_row=len(sectors) + 1,
    )
    sector_chart.add_data(
        sector_data, titles_from_data=True
    )
    sector_chart.set_categories(sector_categories)
    dashboard.add_chart(sector_chart, "J9")

    # Dashboard layout
    for column in range(1, 19):
        dashboard.column_dimensions[
            get_column_letter(column)
        ].width = 15

    dashboard.freeze_panes = "A4"

    # Save only after building the workbook successfully
    wb.save(OUTPUT_FILE)

    print(f"Dashboard created successfully: {OUTPUT_FILE}")
    print(f"Companies: {len(companies)}")
    print(f"Sectors: {len(sectors)}")


if __name__ == "__main__":
    main()
